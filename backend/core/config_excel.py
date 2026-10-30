"""
Conversion de la configuration des entités (entities.yaml) vers et depuis Excel.

Classeur :
- « Entités » : une ligne par entité (activation, masque, description, contexte, validateur) ;
- « Motifs »  : une ligne par expression régulière, rattachée à son entité ;
- « Paramètres » : seuil de détection.
"""

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

ENTITY_HEADERS = ["Entité", "Activée", "Masque", "Description", "Contexte", "Validateur", "Validation stricte"]
PATTERN_HEADERS = ["Entité", "Nom du motif", "Regex", "Score"]
SETTINGS_HEADERS = ["Paramètre", "Valeur"]
CONTEXT_SEPARATOR = ", "

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="5B5BF0")


class ExcelConfigError(ValueError):
    """Classeur illisible ou incomplet (message destiné à l'utilisateur)."""


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
    for row in rows:
        sheet.append(row)
    # Une regex commençant par « = » ne doit jamais devenir une formule
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for index, width in enumerate(widths):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def to_xlsx(config: dict) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    entity_rows, pattern_rows = [], []
    for name, entity in config.get("entities", {}).items():
        entity_rows.append([
            name,
            "oui" if entity.get("enabled", True) else "non",
            entity.get("mask", f"[{name}]"),
            entity.get("description", ""),
            CONTEXT_SEPARATOR.join(entity.get("context") or []),
            entity.get("validator", ""),
            "" if "validator" not in entity else ("oui" if entity.get("strict_validation", True) else "non"),
        ])
        for pattern in entity.get("patterns") or []:
            pattern_rows.append([name, pattern.get("name", ""), pattern.get("regex", ""), pattern.get("score", 0.5)])

    _sheet(workbook, "Entités", ENTITY_HEADERS, entity_rows, [24, 9, 26, 60, 40, 12, 18])
    _sheet(workbook, "Motifs", PATTERN_HEADERS, pattern_rows, [24, 28, 90, 8])
    _sheet(workbook, "Paramètres", SETTINGS_HEADERS, [["Seuil", config.get("threshold", 0.35)]], [16, 12])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _rows(sheet: Worksheet, headers: list[str]):
    """Lignes non vides sous forme de dict, en vérifiant les en-têtes attendus."""
    found = [str(c.value or "").strip() for c in sheet[1]]
    missing = [h for h in headers if h not in found]
    if missing:
        raise ExcelConfigError(f"Feuille « {sheet.title} » : colonne(s) manquante(s) : {', '.join(missing)}")
    for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        row = {h: values[found.index(h)] if found.index(h) < len(values) else None for h in headers}
        if any(v not in (None, "") for v in row.values()):
            yield number, row


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _flag(value, sheet: str, number: int, column: str) -> bool:
    text = _text(value).casefold()
    if text in ("oui", "o", "yes", "y", "true", "vrai", "1", "x"):
        return True
    if text in ("non", "n", "no", "false", "faux", "0", ""):
        return False
    raise ExcelConfigError(f"{sheet}, ligne {number} : « {value} » n'est pas oui/non ({column})")


def from_xlsx(data: bytes) -> tuple[float | None, dict]:
    """Renvoie (seuil ou None, entités au format entities.yaml)."""
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True)
    except Exception as e:
        raise ExcelConfigError(f"Fichier Excel illisible : {e}") from e
    if "Entités" not in workbook.sheetnames:
        raise ExcelConfigError("Feuille « Entités » introuvable : exportez d'abord la configuration pour obtenir le modèle")

    entities: dict = {}
    for number, row in _rows(workbook["Entités"], ENTITY_HEADERS):
        name = _text(row["Entité"])
        if not name:
            raise ExcelConfigError(f"Entités, ligne {number} : nom d'entité vide")
        if name in entities:
            raise ExcelConfigError(f"Entités, ligne {number} : entité {name} en double")
        entity = {
            "enabled": _flag(row["Activée"], "Entités", number, "Activée"),
            "mask": _text(row["Masque"]) or f"[{name}]",
            "description": _text(row["Description"]),
        }
        context = [w.strip() for w in _text(row["Contexte"]).split(",") if w.strip()]
        if context:
            entity["context"] = context
        if _text(row["Validateur"]):
            entity["validator"] = _text(row["Validateur"])
            if _text(row["Validation stricte"]):
                entity["strict_validation"] = _flag(row["Validation stricte"], "Entités", number, "Validation stricte")
        entities[name] = entity

    if "Motifs" in workbook.sheetnames:
        for number, row in _rows(workbook["Motifs"], PATTERN_HEADERS):
            name = _text(row["Entité"])
            if name not in entities:
                raise ExcelConfigError(f"Motifs, ligne {number} : entité « {name} » absente de la feuille Entités")
            try:
                score = float(row["Score"])
            except (TypeError, ValueError):
                raise ExcelConfigError(f"Motifs, ligne {number} : score invalide « {row['Score']} »") from None
            if not 0 < score <= 1:
                raise ExcelConfigError(f"Motifs, ligne {number} : le score doit être compris entre 0 et 1")
            regex = "" if row["Regex"] is None else str(row["Regex"])
            if not regex:
                raise ExcelConfigError(f"Motifs, ligne {number} : regex vide")
            entities[name].setdefault("patterns", []).append({
                "name": _text(row["Nom du motif"]) or f"{name}_{number}",
                "regex": regex,
                "score": score,
            })

    threshold = None
    if "Paramètres" in workbook.sheetnames:
        for number, row in _rows(workbook["Paramètres"], SETTINGS_HEADERS):
            if _text(row["Paramètre"]).casefold() == "seuil":
                try:
                    threshold = float(row["Valeur"])
                except (TypeError, ValueError):
                    raise ExcelConfigError(f"Paramètres, ligne {number} : seuil invalide") from None
                if not 0 <= threshold <= 1:
                    raise ExcelConfigError("Paramètres : le seuil doit être compris entre 0 et 1")

    return threshold, entities
