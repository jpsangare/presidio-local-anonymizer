"""Import / export Excel des formules d'anonymisation."""

import base64
import io

import yaml
from openpyxl import load_workbook

from core.config_excel import from_xlsx, to_xlsx


def export_workbook(client):
    body = client.get("/api/config/export").get_json()
    assert body["filename"].endswith(".xlsx")
    return load_workbook(io.BytesIO(base64.b64decode(body["content"])))


def import_workbook(client, workbook):
    buffer = io.BytesIO()
    workbook.save(buffer)
    return client.post("/api/config/import", json={"content": base64.b64encode(buffer.getvalue()).decode()})


def saved_entities(temp_config):
    return yaml.safe_load(temp_config.read_text(encoding="utf-8"))["entities"]


def test_round_trip_keeps_every_entity_and_pattern(entities_config):
    threshold, entities = from_xlsx(to_xlsx(entities_config))

    assert threshold == entities_config["threshold"]
    assert set(entities) == set(entities_config["entities"])
    for name, original in entities_config["entities"].items():
        assert entities[name].get("patterns", []) == original.get("patterns", []), name
        assert entities[name].get("context", []) == original.get("context", []), name
        assert entities[name]["enabled"] == original.get("enabled", True), name
        assert entities[name].get("validator") == original.get("validator"), name


def test_regex_starting_with_equals_is_not_a_formula():
    config = {"entities": {"X": {"mask": "[X]", "patterns": [{"name": "P", "regex": "=SUM(A1)", "score": 0.5}]}}}

    _, entities = from_xlsx(to_xlsx(config))

    assert entities["X"]["patterns"][0]["regex"] == "=SUM(A1)"


def test_import_updates_and_adds_entities(client, temp_config):
    workbook = export_workbook(client)
    workbook["Entités"].append(["TEST_DOSSIER", "oui", "[DOSSIER]", "Numéro de dossier", "dossier, réf", None, None])
    workbook["Motifs"].append(["TEST_DOSSIER", "DOSSIER", r"\bDOS-\d{6}\b", 0.7])
    workbook["Paramètres"]["B2"] = 0.4

    response = import_workbook(client, workbook)

    assert response.status_code == 200
    assert response.get_json()["added"] == ["TEST_DOSSIER"]
    saved = yaml.safe_load(temp_config.read_text(encoding="utf-8"))
    assert saved["threshold"] == 0.4
    assert saved["entities"]["TEST_DOSSIER"] == {
        "enabled": True, "mask": "[DOSSIER]", "description": "Numéro de dossier",
        "context": ["dossier", "réf"],
        "patterns": [{"name": "DOSSIER", "regex": r"\bDOS-\d{6}\b", "score": 0.7}],
    }


def test_entities_missing_from_workbook_are_kept(client, temp_config):
    workbook = export_workbook(client)
    sheet = workbook["Entités"]
    sheet.delete_rows(3, sheet.max_row)  # ne garde qu'une entité
    workbook["Motifs"].delete_rows(2, workbook["Motifs"].max_row)

    assert import_workbook(client, workbook).status_code == 200
    assert "EMAIL_ADDRESS" in saved_entities(temp_config)


def test_invalid_regex_is_rejected_and_nothing_is_saved(client, temp_config):
    before = temp_config.read_text(encoding="utf-8")
    workbook = export_workbook(client)
    workbook["Motifs"].append(["EMAIL_ADDRESS", "CASSE", "([", 0.5])

    response = import_workbook(client, workbook)

    assert response.status_code == 400
    assert "EMAIL_ADDRESS" in response.get_json()["message"]
    assert temp_config.read_text(encoding="utf-8") == before


def test_errors_point_to_the_faulty_row(client, temp_config):
    workbook = export_workbook(client)
    workbook["Entités"]["B2"] = "peut-être"

    response = import_workbook(client, workbook)

    assert response.status_code == 400
    assert "ligne 2" in response.get_json()["message"]


def test_unknown_validator_is_rejected(client, temp_config):
    workbook = export_workbook(client)
    workbook["Entités"]["F2"] = "inconnu"

    response = import_workbook(client, workbook)

    assert response.status_code == 400
    assert "luhn" in response.get_json()["message"]


def test_not_an_excel_file_returns_400(client, temp_config):
    content = base64.b64encode(b"pas un classeur").decode()

    assert client.post("/api/config/import", json={"content": content}).status_code == 400
